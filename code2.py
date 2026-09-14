import streamlit as st
import pandas as pd
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
from copy import copy
import string
import io

st.set_page_config(page_title="Automação Caderno de Campo - PEPAC", layout="centered")

st.title("🌾 Gerador Automático de Caderno de Campo")
st.markdown("Ferramenta oficial para processamento de exportações do Pedido Único (iSIP/IFAP) e preenchimento inteligente do Caderno de Campo.")

uploaded_pu = st.file_uploader("1. Carregar ficheiro em bruto do Pedido Único (Excel .xlsx)", type=["xlsx"])
uploaded_template = st.file_uploader("2. Carregar o Template Vazio do Caderno de Campo (.xlsx)", type=["xlsx"])

def duplicate_block_precise(ws, start_row, end_row, num_copies):
    block_height = end_row - start_row + 1
    for c in range(1, num_copies + 1):
        offset = block_height * c
        for merged_range in list(ws.merged_cells.ranges):
            if start_row <= merged_range.min_row <= end_row:
                min_col = get_column_letter(merged_range.min_col)
                max_col = get_column_letter(merged_range.max_col)
                ws.merge_cells(f"{min_col}{merged_range.min_row + offset}:{max_col}{merged_range.max_row + offset}")
        for r in range(start_row, end_row + 1):
            if r in ws.row_dimensions:
                ws.row_dimensions[r + offset].height = ws.row_dimensions[r].height
        for row in range(start_row, end_row + 1):
            for col in range(1, ws.max_column + 1):
                source = ws.cell(row, col)
                target = ws.cell(row + offset, col)
                if type(source).__name__ != 'MergedCell' and type(target).__name__ != 'MergedCell':
                    target.value = source.value
                if source.has_style:
                    target.font = copy(source.font)
                    target.border = copy(source.border)
                    target.fill = copy(source.fill)
                    target.number_format = copy(source.number_format)
                    target.protection = copy(source.protection)
                    target.alignment = copy(source.alignment)
    return [start_row] + [start_row + block_height * c for c in range(1, num_copies + 1)]

def safe_write(ws, row, col, value):
    cell = ws.cell(row=row, column=col)
    if type(cell).__name__ == 'MergedCell':
        for merged_range in ws.merged_cells.ranges:
            if cell.coordinate in merged_range:
                ws.cell(row=merged_range.min_row, column=merged_range.min_col).value = value
                return
    else:
        cell.value = value

def process_caderno(pu_file, template_file):
    excel_file = pd.ExcelFile(pu_file)
    sheets = excel_file.sheet_names
    
    df_culturas = pd.read_excel(excel_file, sheet_name='culturas')
    df_sub = pd.read_excel(excel_file, sheet_name='subparcelas')
    df_maa = pd.read_excel(excel_file, sheet_name='medidasAgroAmbientais') if 'medidasAgroAmbientais' in sheets else pd.DataFrame()
    df_rest = pd.read_excel(excel_file, sheet_name='restricoesTerritoriais') if 'restricoesTerritoriais' in sheets else pd.DataFrame()
    
    df_final = pd.DataFrame()
    df_final['Nº seq. de Parcela'] = df_culturas['pk.parNumSeq']
    df_final['Subparcela'] = df_culturas['pk.spaNumero']
    df_final['Área (ha)'] = df_culturas['culAreSegPil']
    
    # Extração direta da coluna Q do separador culturas (índice 16): 'S' para Sequeiro e 'R' para Regadio
    if df_culturas.shape[1] > 16:
        df_final['Regadio'] = df_culturas.iloc[:, 16].fillna('S').astype(str).str.strip().str.upper()
    else:
        reg_col = next((c for c in df_culturas.columns if any(k in str(c).lower() for k in ['reg', 'sec'])), None)
        if reg_col:
            df_final['Regadio'] = df_culturas[reg_col].fillna('S').astype(str).str.strip().str.upper()
        else:
            df_final['Regadio'] = 'S'
    
    tabela_ifap_culturas = {
        '1': '1 - TRIGO',
        '3': '3 - CENTEIO',
        '4': '4 - CEVADA',
        '5': '5 - AVEIA',
        '6': '6 - MILHO',
        '7': '7 - TRITICALE',
        '8': '8 - SORGO',
        '9': '9 - LINHO',
        '13': '13 - ERVILHA',
        '14': '14 - FAVA',
        '17': '17 - GIRASSOL',
        '18': '18 - SOJA',
        '24': '24 - ARROZ',
        '26': '26 - OUTROS CEREAIS',
        '29': '29 - CANA DE AÇÚCAR',
        '32': '32 - BETERRABA',
        '33': '33 - TOMATE',
        '34': '34 - VINHA',
        '35': '35 - LUPULO',
        '38': '38 - GRÃO DE BICO',
        '42': '42 - TABACO',
        '44': '44 - LUZERNA',
        '46': '46 - TREVO',
        '47': '47 - TREMOÇO',
        '48': '48 - ERVILHACA',
        '59': '59 - ALGODÃO',
        '60': '60 - OUTRAS CULTURAS PERMANENTES',
        '67': '67 - AZEVEM',
        '78': '78 - PIMENTO',
        '81': '81 - PLANTAS AROM., MEDICINAIS E CONDIMENTARES',
        '83': '83 - OLIVAL',
        '84': '84 - POMARES MISTOS DE FRUTOS FRESCOS',
        '85': '85 - FIGO',
        '86': '86 - OUTROS FRUTOS SECOS',
        '89': '89 - POUSIO',
        '90': '90 - OUTRAS HORTÍCOLAS',
        '91': '91 - FLORES E PLANTAS ORNAMENTAIS',
        '93': '93 - PERA',
        '94': '94 - PÊSSEGO',
        '96': '96 - LARANJA',
        '97': '97 - LIMÃO',
        '101': '101 - VIVEIROS',
        '102': '102 - OUTROS FRUTOS SUB-TROPICAIS',
        '103': '103 - BATATA',
        '105': '105 - MAÇÃ',
        '106': '106 - CEREJA',
        '107': '107 - DAMASCO',
        '108': '108 - AMEIXA',
        '109': '109 - AMENDOA',
        '110': '110 - CASTANHA',
        '111': '111 - ALFARROBA',
        '112': '112 - NOZ',
        '114': '114 - VIME',
        '115': '115 - CHA',
        '116': '116 - AVELÃ',
        '117': '117 - OUTROS PEQUENOS FRUTOS',
        '118': '118 - MARMELO',
        '119': '119 - NÊSPERA',
        '124': '124 - KIWI',
        '125': '125 - GALERIA RIPÍCOLA',
        '126': '126 - SALIX (MADEIRA)',
        '127': '127 - BATATA DOCE',
        '128': '128 - INHAME',
        '130': '130 - AMENDOIM',
        '131': '131 - MEDRONHEIRO',
        '132': '132 - ANANÁS',
        '133': '133 - BANANA',
        '134': '134 - PISTACIOS',
        '135': '135 - PINHÃO',
        '136': '136 - ABACATE',
        '137': '137 - BERINGELA',
        '139': '139 - BOSQUETES',
        '142': '142 - PRADOS TEMPORÁRIOS',
        '143': '143 - PASTAGENS PERMANENTES',
        '148': '148 - OUTRAS LEGUMINOSAS SECAS',
        '151': '151 - ANONA',
        '155': '155 - TANGERINA',
        '156': '156 - OUTRAS OLEAGINOSAS',
        '157': '157 - OUTROS CITRINOS',
        '161': '161 - MISTO CULTURAS PERMANENTES',
        '162': '162 - POVOAMENTO DE SOBREIROS',
        '163': '163 - POVOAMENTO AZINHEIRAS',
        '164': '164 - POVOAMENTO CARVALHO NEGRAL',
        '165': '165 - POVOAMENTO MISTO QUERCUS(SOB.AZENH.CARVAL/NEGRAL)',
        '166': '166 - POVOAMENTO CASTANHEIRO',
        '167': '167 - POVOAMENTO OUTRAS FOLHOSAS',
        '168': '168 - POVOAMENTO DE PINHEIRO MANSO',
        '169': '169 - POVOAMENTO OUTRAS RESINOSAS',
        '170': '170 - POVOAMENTO F MISTO',
        '173': '173 - ACEIRO FLORESTAL',
        '174': '174 - OUTRAS SUPERFÍCIES FLORESTAIS',
        '190': '190 - MACIÇOS OU FORMAÇÕES RELIQUIAIS OU NOTÁVEIS',
        '195': '195 - OUTRAS FRUTOS FRESCOS',
        '201': '201 - AMORA',
        '202': '202 - MIRTILO',
        '203': '203 - FRAMBOESA',
        '204': '204 - MORANGO',
        '205': '205 - MELÃO',
        '208': '208 - DIOSPIRO',
        '209': '209 - ROMÃ',
        '210': '210 - POVOAMENTO DE EUCALIPTO',
        '211': '211 - GINJA',
        '223': '223 - SABUGUEIRO (BAGA)',
        '230': '230 - FEIJÃO',
        '231': '231 - MELANCIA',
        '232': '232 - MELOA',
        '233': '233 - NABO',
        '234': '234 - PEPINO',
        '236': '236 - RABANETE',
        '237': '237 - RÁBANO',
        '238': '238 - RUTABAGA',
        '240': '240 - TREMOCILHA',
        '241': '241 - ABÓBORAS E ABOBORINHAS',
        '242': '242 - AGRIÃO',
        '244': '244 - ALFACE',
        '245': '245 - ALHO',
        '248': '248 - CEBOLA',
        '249': '249 - CENOURA',
        '250': '250 - COURGETTE',
        '254': '254 - COUVE',
        '261': '261 - TANGERA',
        '262': '262 - SOBREIRO PARA PRODUÇÃO DE CORTIÇA',
        '263': '263 - CHUCHU',
        '264': '264 - COLZA',
        '265': '265 - SUPERFÍCIE ARBUSTIVA NÃO PASTOREÁVEL',
        '266': '266 - CONSOCIAÇÃO DE FIXADORAS DE AZOTO',
        '267': '267 - CONSOCIAÇÕES ANUAIS E OUTRAS CULT. FORRAG. ANUAIS',
        '268': '268 - GALERIA RIPÍCOLA FLORESTAL',
        '269': '269 - MARACUJÁ',
        '276': '276 - MOSTARDA',
        '277': '277 - NABIÇA',
        '279': '279 - RÚCULA',
        '280': '280 - PASTAGENS ARBUSTIVAS',
        '281': '281 - PASTAGENS EM BALDIO',
        '283': '283 - MEDRONHO',
        '284': '284 - PAPAIA',
        '285': '285 - FIGO DA INDIA',
        '286': '286 - GROSELHA',
        '287': '287 - SERRADELA',
        '288': '288 - FESTUCA',
        '289': '289 - PANASCO',
        '290': '290 - BROMUS',
        '291': '291 - MARALFALFA',
        '292': '292 - ANAFA',
        '293': '293 - ALHO FRANCÊS',
        '298': '298 - MANGA',
        '299': '299 - GOIABA',
        '300': '300 - TALUDE DA VINHA',
        '301': '301 - CULTURAS EM HIDROPONIA',
        '302': '302 - CULTURAS SEM SOLO',
        '304': '304 - TRIGO SPELTA',
        '305': '305 - ESPINAFRE',
        '914': '914 - ELEMENTO LINEAR SEBE OU CORTA-VENTO-ÁREA ÚTIL',
        '924': '924 - ELEMENTO LINEAR EM ORIZICULTURA-ÁREA ÚTIL',
        '925': '925 - GALERIA RIPÍCOLA - ÁREA ÚTIL',
        '939': '939 - EP-BOSQUETE E FORMAÇÕES RELIQUIAIS-ÁREA ÚTIL',
        '982': '982 - CABECEIRAS CULT. PERMANENTES -ÁREA ÚTIL',
        '985': '985 - LINHAS DE ÁGUA - ÁREA ÚTIL',
        '986': '986 - ELP CHARCAS E LAGOAS - ÁREA ÚTIL',
    }

    def map_cultura(code):
        code_str = str(code).strip()
        if code_str in tabela_ifap_culturas:
            return tabela_ifap_culturas[code_str]
        try:
            c_int = str(int(code_str))
            if c_int in tabela_ifap_culturas:
                return tabela_ifap_culturas[c_int]
        except:
            pass
        return f"{code_str} - Cultura IFAP"

    df_final['Cultura/Variedade ou casta'] = df_culturas['culCodigo'].apply(map_cultura)
    df_final['Textura do solo'] = ""
    df_final['Sucessão cultural'] = "-"
    df_final['Boas práticas'] = ""

    df_final = pd.merge(df_final, df_sub[['pk.parNumSeq', 'pk.spaNumero', 'spaIqfp']], 
                        left_on=['Nº seq. de Parcela', 'Subparcela'], 
                        right_on=['pk.parNumSeq', 'pk.spaNumero'], how='left')
    df_final['IQFP'] = df_final['spaIqfp']
    df_final = df_final.drop(columns=['pk.parNumSeq', 'pk.spaNumero', 'spaIqfp'])

    # Dicionário estrito de Intervenções e Medidas do PEPAC (Apenas Abreviaturas)
    code_mapping = {
        # C.1.1.1 - Uso Eficiente dos Recursos Naturais e submedidas
        'C.1.1.1.1.1': 'SD',
        'C.1.1.1.1.2': 'ENR',
        'C.1.1.1.1.3': 'PB',
        'C.1.1.1.2': 'UEA-Água',
        
        # C.1.1.2 - Manutenção de Sistemas Extensivos
        'C.1.1.2': 'MSE',
        'C.1.1.2.1': 'ML',
        'C.1.1.2.2': 'CPPT',
        'C.1.1.2.2.1.1': 'OLI-TRAD',
        'C.1.1.2.2.1.2': 'FIG-EXT',
        'C.1.1.2.2.1.3': 'POM-TRAD',
        'C.1.1.2.2.1.4': 'AMEND-EXT',
        'C.1.1.2.2.1.5': 'CAST-EXT',
        'C.1.1.2.3': 'MA',
        
        # C.1.1.3 a C.1.1.6
        'C.1.1.3': 'CB',
        'C.1.1.4': 'GAPB',
        'C.1.1.5': 'CMRGF',
        'C.1.1.6': 'APIO',
        
        # C.1.1.7 (PRODI) e C.1.1.8 (AB / AB-C)
        'C.1.1.7': 'PRODI',
        'C.1.1.8': 'AB',
        'C.1.1.8.1': 'AB-C',
        'C.1.1.8.2': 'AB',
        
        # C.1.2 - Condicionantes e Rede Natura
        'C.1.2.1': 'ZM',
        'C.1.2.2': 'RN',
        'C.1.2.3': 'ZCNE',
        
        # Eixo A - Ecorregimes
        'A.3.3': 'ECO-BIO',
        'A.3.4': 'MSPP',
        'A.3.5': 'MEAF',
        'A.3.6': 'CERT',
        
        # Eixo D - Planos Zonais e Específicas
        'D.2.1': 'PZA',
        'D.2.4': 'PE-AVES',
        'D.2.5': 'PE-SILVO'
    }

    intervencoes_lista = []
    if not df_maa.empty and 'pk.parNumSeq' in df_maa.columns and 'pk.spaNumero' in df_maa.columns and 'intCodigo' in df_maa.columns:
        for _, r in df_maa.iterrows():
            p = r.get('pk.parNumSeq')
            s = r.get('pk.spaNumero')
            val = str(r.get('intCodigo', '')).strip()
            if val and val != 'nan':
                val_traduzido = code_mapping.get(val, val)
                intervencoes_lista.append({'pk.parNumSeq': p, 'pk.spaNumero': s, 'interv': val_traduzido})

    if not df_rest.empty and 'pk.parNumSeq' in df_rest.columns and 'pk.spaNumero' in df_rest.columns:
        for col_r in df_rest.columns:
            if 'cod' in col_r.lower() or 'tipo' in col_r.lower():
                for _, r in df_rest.iterrows():
                    p = r.get('pk.parNumSeq')
                    s = r.get('pk.spaNumero')
                    val = str(r.get(col_r, '')).strip()
                    if val and val != 'nan':
                        val_traduzido = code_mapping.get(val, val)
                        intervencoes_lista.append({'pk.parNumSeq': p, 'pk.spaNumero': s, 'interv': val_traduzido})

    if intervencoes_lista:
        df_int_all = pd.DataFrame(intervencoes_lista).drop_duplicates()
        grouped_int = df_int_all.groupby(['pk.parNumSeq', 'pk.spaNumero'])['interv'].apply(lambda x: ', '.join(x.dropna().astype(str))).reset_index()
        grouped_int.columns = ['Nº seq. de Parcela', 'Subparcela', 'Intervenções_Todas']
        df_final = pd.merge(df_final, grouped_int, on=['Nº seq. de Parcela', 'Subparcela'], how='left')
        df_final['Intervenção PEPAC'] = df_final['Intervenções_Todas'].fillna("-")
        df_final = df_final.drop(columns=['Intervenções_Todas'], errors='ignore')
    else:
        df_final['Intervenção PEPAC'] = "-"

    df_final['Modo de Produção'] = ""
    df_final['Zona Homogénea'] = ""

    culture_modes = {}
    for idx, row in df_final.iterrows():
        cult = row['Cultura/Variedade ou casta']
        val = row['Intervenção PEPAC']
        if pd.notna(cult):
            if cult not in culture_modes:
                culture_modes[cult] = {'AB': False, 'PRODI': False}
            if pd.notna(val):
                val_upper = str(val).upper()
                if "AB" in val_upper or "BIOL" in val_upper: culture_modes[cult]['AB'] = True
                if "PRODI" in val_upper: culture_modes[cult]['PRODI'] = True

    for idx, row in df_final.iterrows():
        cult = row['Cultura/Variedade ou casta']
        if pd.notna(cult):
            modes = []
            if culture_modes[cult]['AB']: modes.append("AB")
            if culture_modes[cult]['PRODI']: modes.append("PRODI")
            df_final.at[idx, 'Modo de Produção'] = ", ".join(modes) if modes else "CV"

    unique_cultures = list(df_final['Cultura/Variedade ou casta'].dropna().unique())
    def is_priority(c):
        return culture_modes.get(c, {}).get('AB', False) or culture_modes.get(c, {}).get('PRODI', False)

    priority_cultures = [c for c in unique_cultures if is_priority(c)]
    cv_cultures = [c for c in unique_cultures if not is_priority(c)]

    def sort_group(cult_list):
        vinha = [c for c in cult_list if "034" in c or "Vinha" in c]
        cab = [c for c in cult_list if "982" in c or "Cabeceiras" in c]
        others = [c for c in cult_list if c not in vinha and c not in cab]
        return vinha + cab + others

    sorted_cultures = sort_group(priority_cultures) + sort_group(cv_cultures)
    alphabet = list(string.ascii_uppercase)
    culture_mapping = {}
    for i, culture in enumerate(sorted_cultures):
        letter = alphabet[i] if i < 26 else alphabet[(i//26)-1] + alphabet[i%26]
        culture_mapping[culture] = letter

    for idx, row in df_final.iterrows():
        cult = row['Cultura/Variedade ou casta']
        if pd.notna(cult) and cult in culture_mapping:
            df_final.at[idx, 'Zona Homogénea'] = culture_mapping[cult]

    cols = ['Nº seq. de Parcela', 'Subparcela', 'Zona Homogénea', 'Modo de Produção', 'Intervenção PEPAC', 
            'Área (ha)', 'Textura do solo', 'Cultura/Variedade ou casta', 'Sucessão cultural', 'IQFP', 'Boas práticas']

    wb = load_workbook(template_file)
    ws2 = wb['2 Caraterização área sob comp.']
    for row_idx, row_data in enumerate(df_final[cols].values, start=8):
        for col_idx, value in enumerate(row_data, start=1):
            safe_write(ws2, row_idx, col_idx, value)

    zonas = df_final.groupby('Zona Homogénea').agg({
        'Área (ha)': 'sum',
        'Cultura/Variedade ou casta': 'first',
        'Modo de Produção': 'first',
        'Regadio': lambda x: 'R' if 'R' in list(x) else 'S'
    }).reset_index()

    ws5a = wb['5A Registo Operações Fertil']
    has_prodi = any("PRODI" in str(m) for m in df_final['Modo de Produção'])
    has_ab = any("AB" in str(m) for m in df_final['Modo de Produção'])
    if has_prodi or has_ab:
        for i, row in enumerate(zonas.values, start=5):
            safe_write(ws5a, i, 1, f"{row[0]} ({row[1]})")
            safe_write(ws5a, i, 7, row[2])

    def get_rega(cultura, regadio='S'):
        reg = str(regadio).strip().upper()
        if reg == 'S':
            return "Não"
        elif reg == 'R':
            cult_str = str(cultura).lower()
            if any(x in cult_str for x in ["amendoal", "olival", "vinha", "noz", "tomate", "melão", "citrinos", "laranja", "pomar"]):
                return "gota-a-gota"
            return "Sim"
        return "Não"

    def get_esperada(cultura):
        cult_str = str(cultura).lower()
        if any(x in cult_str for x in ["cabeceiras", "vala", "ripícola", "bosquete", "pousio", "charca", "elp", "lagoas"]):
            return ""
        if any(x in cult_str for x in ["pastagem", "pastagens", "prado", "prados"]):
            return "2"
        if "noz" in cult_str: return "3,5" 
        if "amendoal" in cult_str: return "1,5"
        if "olival" in cult_str: return "3"
        if "vinha" in cult_str: return "6"
        if "melão" in cult_str: return "25"
        if "tomate" in cult_str: return "85"
        if "figo" in cult_str: return "15" 
        return "2"

    ws4 = wb['4 Registo Prot Fitossanitária']
    num_zonas_4 = len(zonas)
    if num_zonas_4 > 1: offsets_4 = duplicate_block_precise(ws4, 4, 27, num_zonas_4 - 1)
    else: offsets_4 = [4]

    for i, row in zonas.iterrows():
        offset = offsets_4[i] - 4
        cult = row['Cultura/Variedade ou casta']
        esp = get_esperada(cult)
        
        safe_write(ws4, offset+5, 2, row['Zona Homogénea'])
        safe_write(ws4, offset+5, 4, round(row['Área (ha)'], 2))
        safe_write(ws4, offset+5, 7, get_rega(cult, row.get('Regadio', 'S'))) 
        safe_write(ws4, offset+7, 3, cult)
        if esp != "":
            safe_write(ws4, offset+11, 4, esp)

    ws5 = wb['5 Registo Operações Culturais ']
    zonas_5 = zonas[zonas['Modo de Produção'].astype(str).str.contains('AB|PRODI', na=False)].reset_index(drop=True)
    num_zonas_5 = len(zonas_5)
    if num_zonas_5 > 0:
        if num_zonas_5 > 1: offsets_5 = duplicate_block_precise(ws5, 4, 31, num_zonas_5 - 1)
        else: offsets_5 = [4]
            
        for i, row in zonas_5.iterrows():
            offset = offsets_5[i] - 4
            cult = row['Cultura/Variedade ou casta']
            esp = get_esperada(cult)
            
            safe_write(ws5, offset+5, 2, row['Zona Homogénea']) 
            safe_write(ws5, offset+5, 12, round(row['Área (ha)'], 2)) 
            safe_write(ws5, offset+7, 2, cult) 
            safe_write(ws5, offset+9, 14, get_rega(cult, row.get('Regadio', 'S'))) 
            if esp != "":
                safe_write(ws5, offset+12, 4, esp) 

    output_buffer = io.BytesIO()
    wb.save(output_buffer)
    output_buffer.seek(0)
    return output_buffer

if uploaded_pu and uploaded_template:
    if st.button("🚀 Processar e Gerar Caderno de Campo", type="primary"):
        with st.spinner("A processar dados do Pedido Único e a estruturar o Caderno de Campo..."):
            try:
                result_stream = process_caderno(uploaded_pu, uploaded_template)
                st.success("Caderno de Campo gerado com sucesso!")
                st.download_button(
                    label="📥 Descarregar Caderno de Campo Preenchido",
                    data=result_stream,
                    file_name="Caderno_de_Campo_Final.xlsx",
                    mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
                )
            except Exception as e:
                st.error(f"Ocorreu um erro durante o processamento: {e}")
else:
    st.info("👆 Por favor, carrega o ficheiro do Pedido Único e o template vazio do Caderno de Campo para começar.")