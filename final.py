import streamlit as st
import pandas as pd
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
from copy import copy
import string
import io

st.set_page_config(page_title="Automação Caderno de Campo - PEPAC", layout="centered")

st.title("🌾 Gerador Automático de Caderno de Campo")
st.markdown("Ferramenta oficial para processamento de exportações do Pedido Único (iSIP/IFAP) e preenchimento inteligente do Caderno de Campo.")

uploaded_pu = st.file_uploader("1. Carregar ficheiro em bruto do Pedido Único (Excel .xlsx)", type=["xlsx"])
uploaded_template = st.file_uploader("2. Carregar o Template Vazio do Caderno de Campo (.xlsx)", type=["xlsx"])

def duplicate_block_precise(ws, start_row, end_row, num_copies):
    block_height = end_row - start_row + 1
    for c in range(1, num_copies + 1):
        offset = block_height * c
        for merged_range in list(ws.merged_cells.ranges):
            if start_row <= merged_range.min_row <= end_row:
                min_col = get_column_letter(merged_range.min_col)
                max_col = get_column_letter(merged_range.max_col)
                ws.merge_cells(f"{min_col}{merged_range.min_row + offset}:{max_col}{merged_range.max_row + offset}")
        for r in range(start_row, end_row + 1):
            if r in ws.row_dimensions:
                ws.row_dimensions[r + offset].height = ws.row_dimensions[r].height
        for row in range(start_row, end_row + 1):
            for col in range(1, ws.max_column + 1):
                source = ws.cell(row, col)
                target = ws.cell(row + offset, col)
                if type(source).__name__ != 'MergedCell' and type(target).__name__ != 'MergedCell':
                    target.value = source.value
                if source.has_style:
                    target.font = copy(source.font)
                    target.border = copy(source.border)
                    target.fill = copy(source.fill)
                    target.number_format = copy(source.number_format)
                    target.protection = copy(source.protection)
                    target.alignment = copy(source.alignment)
    return [start_row] + [start_row + block_height * c for c in range(1, num_copies + 1)]

def safe_write(ws, row, col, value):
    cell = ws.cell(row=row, column=col)
    if type(cell).__name__ == 'MergedCell':
        for merged_range in ws.merged_cells.ranges:
            if cell.coordinate in merged_range:
                ws.cell(row=merged_range.min_row, column=merged_range.min_col).value = value
                return
    else:
        cell.value = value

def process_caderno(pu_file, template_file):
    df_culturas = pd.read_excel(pu_file, sheet_name='culturas')
    df_sub = pd.read_excel(pu_file, sheet_name='subparcelas')
    df_maa = pd.read_excel(pu_file, sheet_name='medidasAgroAmbientais')
    
    df_final = pd.DataFrame()
    df_final['Nº seq. de Parcela'] = df_culturas['pk.parNumSeq']
    df_final['Subparcela'] = df_culturas['pk.spaNumero']
    df_final['Área (ha)'] = df_culturas['culAreSegPil'] # Segundo Pilar
    
    # Leitura universal: Junta o código oficial e a descrição real que vem no ficheiro do IFAP
    def get_nome_cultura(row):
        codigo = str(row.get('culCodigo', ''))
        # Procura por colunas comuns de descrição no export do IFAP (ex: culDescricao, desCultura, nome, etc.)
        descricao = ""
        for col in df_culturas.columns:
            if any(k in col.lower() for k in ['desc', 'nome', 'cultura', 'design']):
                val = str(row.get(col, ''))
                if val and val != 'nan' and val != codigo:
                    descricao = val
                    break
        
        if descricao:
            return f"{codigo} - {descricao}"
        return f"{codigo} - [Cultura IFAP]"

    df_final['Cultura/Variedade ou casta'] = df_culturas.apply(get_nome_cultura, axis=1)
    df_final['Textura do solo'] = ""
    df_final['Sucessão cultural'] = "-"
    df_final['Boas práticas'] = "" 

    df_final = pd.merge(df_final, df_sub[['pk.parNumSeq', 'pk.spaNumero', 'spaIqfp']], 
                        left_on=['Nº seq. de Parcela', 'Subparcela'], 
                        right_on=['pk.parNumSeq', 'pk.spaNumero'], how='left')
    df_final['IQFP'] = df_final['spaIqfp']
    df_final = df_final.drop(columns=['pk.parNumSeq', 'pk.spaNumero', 'spaIqfp'])

    # Tratamento universal de medidas / intervenções agroambientais
    def get_intervencao(row):
        for col in df_maa.columns:
            if any(k in col.lower() for k in ['cod', 'int', 'medida', 'desc']):
                val = str(row.get(col, ''))
                if val and val != 'nan':
                    return val
        return "-"

    if not df_maa.empty and 'pk.parNumSeq' in df_maa.columns and 'pk.spaNumero' in df_maa.columns:
        df_maa['interv_texto'] = df_maa.apply(get_intervencao, axis=1)
        grouped_maa = df_maa.groupby(['pk.parNumSeq', 'pk.spaNumero'])['interv_texto'].apply(lambda x: ', '.join(x.dropna().astype(str))).reset_index()
        grouped_maa.columns = ['Nº seq. de Parcela', 'Subparcela', 'Intervenções']
        df_final = pd.merge(df_final, grouped_maa, on=['Nº seq. de Parcela', 'Subparcela'], how='left')
        df_final['Intervenção PEPAC'] = df_final['Intervenções'].fillna("-")
        df_final = df_final.drop(columns=['Intervenções'], errors='ignore')
    else:
        df_final['Intervenção PEPAC'] = "-"

    df_final['Modo de Produção'] = ""
    df_final['Zona Homogénea'] = ""

    culture_modes = {}
    for idx, row in df_final.iterrows():
        cult = row['Cultura/Variedade ou casta']
        val = row['Intervenção PEPAC']
        if pd.notna(cult):
            if cult not in culture_modes:
                culture_modes[cult] = {'AB': False, 'PRODI': False}
            if pd.notna(val):
                if "AB" in str(val).upper() or "BIOL" in str(val).upper(): culture_modes[cult]['AB'] = True
                if "PRODI" in str(val).upper(): culture_modes[cult]['PRODI'] = True

    for idx, row in df_final.iterrows():
        cult = row['Cultura/Variedade ou casta']
        if pd.notna(cult):
            modes = []
            if culture_modes[cult]['AB']: modes.append("AB")
            if culture_modes[cult]['PRODI']: modes.append("PRODI")
            df_final.at[idx, 'Modo de Produção'] = ", ".join(modes) if modes else "CV"

    unique_cultures = list(df_final['Cultura/Variedade ou casta'].dropna().unique())
    def is_priority(c):
        return culture_modes.get(c, {}).get('AB', False) or culture_modes.get(c, {}).get('PRODI', False)

    priority_cultures = [c for c in unique_cultures if is_priority(c)]
    cv_cultures = [c for c in unique_cultures if not is_priority(c)]

    def sort_group(cult_list):
        vinha = [c for c in cult_list if "034" in c or "vinha" in c.lower()]
        cab = [c for c in cult_list if "982" in c or "cabeceira" in c.lower()]
        others = [c for c in cult_list if c not in vinha and c not in cab]
        return vinha + cab + others

    sorted_cultures = sort_group(priority_cultures) + sort_group(cv_cultures)
    alphabet = list(string.ascii_uppercase)
    culture_mapping = {}
    for i, culture in enumerate(sorted_cultures):
        letter = alphabet[i] if i < 26 else alphabet[(i//26)-1] + alphabet[i%26]
        culture_mapping[culture] = letter

    for idx, row in df_final.iterrows():
        cult = row['Cultura/Variedade ou casta']
        if pd.notna(cult) and cult in culture_mapping:
            df_final.at[idx, 'Zona Homogénea'] = culture_mapping[cult]

    cols = ['Nº seq. de Parcela', 'Subparcela', 'Zona Homogénea', 'Modo de Produção', 'Intervenção PEPAC', 
            'Área (ha)', 'Textura do solo', 'Cultura/Variedade ou casta', 'Sucessão cultural', 'IQFP', 'Boas práticas']
    df_final = df_final[cols]

    wb = load_workbook(template_file)

    ws2 = wb['2 Caraterização área sob comp.']
    for row_idx, row_data in enumerate(df_final.values, start=8):
        for col_idx, value in enumerate(row_data, start=1):
            safe_write(ws2, row_idx, col_idx, value)

    zonas = df_final.groupby('Zona Homogénea').agg({
        'Área (ha)': 'sum',
        'Cultura/Variedade ou casta': 'first',
        'Modo de Produção': 'first'
    }).reset_index()

    ws5a = wb['5A Registo Operações Fertil']
    has_prodi = any("PRODI" in str(m) for m in df_final['Modo de Produção'])
    has_ab = any("AB" in str(m) for m in df_final['Modo de Produção'])
    if has_prodi or has_ab:
        for i, row in enumerate(zonas.values, start=5):
            safe_write(ws5a, i, 1, f"{row[0]} ({row[1]})")
            safe_write(ws5a, i, 7, row[2])

    def get_rega(cultura):
        cult_str = str(cultura).lower()
        if any(x in cult_str for x in ["amendoal", "olival", "azeit", "vinha", "noz", "tomate", "melão", "regado"]):
            return "gota-a-gota"
        return "Não"

    def get_esperada(cultura):
        cult_str = str(cultura).lower()
        if any(x in cult_str for x in ["cabeceiras", "vala", "ripícola", "bosquete", "pousio", "charca", "drenagem"]):
            return ""
        if any(x in cult_str for x in ["pastagem", "pastagens", "prado", "prados"]):
            return "2"
        if "noz" in cult_str: return "3,5" 
        if "amendoal" in cult_str: return "1,5"
        if "olival" in cult_str or "azeit" in cult_str: return "3"
        if "vinha" in cult_str: return "6"
        if "melão" in cult_str: return "25"
        if "tomate" in cult_str: return "85"
        if "trigo" in cult_str or "cevada" in cult_str: return "3,5"
        return "2"

    ws4 = wb['4 Registo Prot Fitossanitária']
    num_zonas_4 = len(zonas)
    if num_zonas_4 > 1: offsets_4 = duplicate_block_precise(ws4, 4, 27, num_zonas_4 - 1)
    else: offsets_4 = [4]

    for i, row in zonas.iterrows():
        offset = offsets_4[i] - 4
        cult = row['Cultura/Variedade ou casta']
        esp = get_esperada(cult)
        
        safe_write(ws4, offset+5, 2, row['Zona Homogénea'])
        safe_write(ws4, offset+5, 4, round(row['Área (ha)'], 2))
        safe_write(ws4, offset+5, 7, get_rega(cult)) 
        safe_write(ws4, offset+7, 3, cult)
        if esp != "":
            safe_write(ws4, offset+11, 4, esp)

    ws5 = wb['5 Registo Operações Culturais ']
    zonas_5 = zonas[zonas['Modo de Produção'].astype(str).str.contains('AB|PRODI', na=False)].reset_index(drop=True)
    num_zonas_5 = len(zonas_5)
    if num_zonas_5 > 0:
        if num_zonas_5 > 1: offsets_5 = duplicate_block_precise(ws5, 4, 31, num_zonas_5 - 1)
        else: offsets_5 = [4]
            
        for i, row in zonas_5.iterrows():
            offset = offsets_5[i] - 4
            cult = row['Cultura/Variedade ou casta']
            esp = get_esperada(cult)
            
            safe_write(ws5, offset+5, 2, row['Zona Homogénea']) 
            safe_write(ws5, offset+5, 12, round(row['Área (ha)'], 2)) 
            safe_write(ws5, offset+7, 2, cult) 
            safe_write(ws5, offset+9, 14, get_rega(cult)) 
            if esp != "":
                safe_write(ws5, offset+12, 4, esp) 

    try:
        df_efetivos = pd.read_excel(pu_file, sheet_name='efetivosPecuariosProprio')
        if not df_efetivos.empty:
            ws3 = wb['3 Caraterização do Efe.Pecuário']
            row_idx_efet = 6
            for _, efetivo in df_efetivos.iterrows():
                if 'eceCodigo' in efetivo and 'eppQtdPas' in efetivo:
                    safe_write(ws3, row_idx_efet, 2, str(efetivo['eceCodigo']))
                    safe_write(ws3, row_idx_efet, 4, str(efetivo['eppQtdPas']))
                    row_idx_efet += 1
    except ValueError:
        pass
        
    output_buffer = io.BytesIO()
    wb.save(output_buffer)
    output_buffer.seek(0)
    return output_buffer

if uploaded_pu and uploaded_template:
    if st.button("🚀 Processar e Gerar Caderno de Campo", type="primary"):
        with st.spinner("A processar dados do Pedido Único e a estruturar o Caderno de Campo..."):
            try:
                result_stream = process_caderno(uploaded_pu, uploaded_template)
                st.success("Caderno de Campo gerado com sucesso!")
                st.download_button(
                    label="📥 Descarregar Caderno de Campo Preenchido",
                    data=result_stream,
                    file_name="Caderno_de_Campo_Final.xlsx",
                    mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
                )
            except Exception as e:
                st.error(f"Ocorreu um erro durante o processamento: {e}")
else:
    st.info("👆 Por favor, carrega o ficheiro do Pedido Único e o template vazio do Caderno de Campo para começar.")